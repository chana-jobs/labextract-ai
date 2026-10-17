"""
Official Excel export (Spec v3 sections 8.3, 11.1, 11.2; P1 1.1, 1.3, 2.2).

One sheet, seven columns in the challenge order, one row per commitment.
Unapproved cell -> empty + orange. Confirmed-absent -> empty, no colour.
Every cell that comes from a document is written as a TEXT string (leading zeros kept,
and a value starting with = + - @ can never become a formula).
"""
from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from pydantic import BaseModel

from shared.contracts import (
    CommitmentStatus, ExportCell, Form, IssueReason, ProcessingStatus, derive_commitment_status,
    export_cell, export_central_cell,
)

HEADERS = ["שם פרטי", "שם משפחה", "מספר ת״ז", "מספר מרכזי", "תאריך קבלת הדגימה", "גורם שולח", "מספר התחייבות"]
ORANGE = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
SHEET_NAME = "נתונים"


class ExportSummary(BaseModel):
    rows: int = 0
    ready: int = 0
    closed_with_gaps: int = 0
    pending_rows: int = 0            # rows with at least one orange cell
    orange_cells: int = 0
    uncounted_forms: int = 0         # COMMITMENT_COUNT_UNCERTAIN: rows may be missing (P1 1.1)
    failed_forms: int = 0
    forms_without_commitments: int = 0
    filename: str = "labextract.xlsx"
    is_partial: bool = False
    warnings: list[str] = []


def _display_date(cell: ExportCell) -> ExportCell:
    """Stored as ISO (YYYY-MM-DD); shown as DD/MM/YYYY. ASSUMPTION: text cell until the lab template is known."""
    if cell.value:
        try:
            return ExportCell(value=datetime.strptime(cell.value, "%Y-%m-%d").strftime("%d/%m/%Y"), orange=False)
        except ValueError:
            return ExportCell(value=None, orange=True)   # unparseable normalized date: never guess
    return cell


def _put(ws, row: int, col: int, cell: ExportCell) -> None:
    c = ws.cell(row=row, column=col)
    c.value = cell.value if cell.value is not None else None
    if cell.value is not None:
        c.data_type = "s"            # force string: no formula injection, zeros preserved
        c.number_format = "@"
    if cell.orange:
        c.fill = ORANGE
    c.alignment = Alignment(horizontal="right")


def build_export(forms: list[Form]) -> tuple[bytes, ExportSummary]:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.sheet_view.rightToLeft = True
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="right")

    s = ExportSummary()
    row = 2
    for form in forms:
        if form.processing_status == ProcessingStatus.FAILED:
            s.failed_forms += 1
            continue
        if any(i.reason == IssueReason.COMMITMENT_COUNT_UNCERTAIN for i in form.issues):
            s.uncounted_forms += 1
        for exam in form.examinations:
            if not exam.commitments:
                s.forms_without_commitments += 1     # no invented row (spec 11.2)
                continue
            shared = [
                export_cell(exam.first_name), export_cell(exam.last_name), export_cell(exam.israeli_id),
                export_central_cell(exam.central_numbers), _display_date(export_cell(exam.sample_received_date)),
                export_cell(exam.sender),
            ]
            for com in exam.commitments:
                cells = shared + [export_cell(com.number)]
                for col, cell in enumerate(cells, start=1):
                    _put(ws, row, col, cell)
                orange = sum(1 for c in cells if c.orange)
                s.orange_cells += orange
                s.pending_rows += 1 if orange else 0
                status = derive_commitment_status(exam, com)
                s.ready += status == CommitmentStatus.READY
                s.closed_with_gaps += status == CommitmentStatus.READY_WITH_CONFIRMED_GAPS
                s.rows += 1
                row += 1

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{max(row - 1, 1)}"
    for col, w in zip("ABCDEFG", (14, 16, 14, 22, 20, 16, 18)):
        ws.column_dimensions[col].width = w

    s.is_partial = bool(s.orange_cells or s.uncounted_forms or s.failed_forms)
    s.filename = "labextract" + ("_חלקי" if s.is_partial else "") + ".xlsx"
    if s.orange_cells:
        s.warnings.append(f"{s.orange_cells} תאים לא מאושרים (ריקים וכתומים) ב-{s.pending_rows} שורות.")
    if s.uncounted_forms:
        s.warnings.append(f"{s.uncounted_forms} טפסים שלא ניתן לקבוע בהם כמה התחייבויות יש: ייתכן שחסרות שורות בקובץ.")
    if s.failed_forms:
        s.warnings.append(f"{s.failed_forms} טפסים נכשלו בעיבוד טכני ולא נכללו בקובץ.")
    if s.is_partial:
        s.warnings.append("הקובץ אינו מוכן לקליטה אוטומטית עד לאימות מול תבנית המעבדה. כתום הוא סימון חזותי בלבד.")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), s
