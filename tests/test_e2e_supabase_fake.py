"""The cloud code path end to end: real sign-in, SupabaseRepository (lx_* tables, private bucket)
against the in-memory PostgREST/Storage fake, through the /api mount used on Vercel.
No network; synthetic users and documents only."""
import io

import pymupdf
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend import auth as auth_module
from backend.app import app as api
from backend.asgi import app
from backend.auth import LocalAuthenticator, LoginRateLimiter, get_authenticator, hash_password, token_hash
from backend.storage import get_repository
from backend.supabase_repo import SupabaseRepository
from tests.fake_supabase import FakeSupabase

pytestmark = pytest.mark.real_auth

USER_A, USER_B = "aaaaaaaa-0000-4000-8000-00000000000a", "bbbbbbbb-0000-4000-8000-00000000000b"
EMAIL_A, EMAIL_B, PASSWORD = "secretary.a@synthetic.test", "secretary.b@synthetic.test", "synthetic-Passw0rd"
BASE = "https://labextract-ai-preview.vercel.app"


@pytest.fixture
def fake(monkeypatch):
    fake = FakeSupabase()
    fake.tables["lx_workspace_members"] += [{"user_id": USER_A, "workspace_id": "demo"},
                                            {"user_id": USER_B, "workspace_id": "demo-b"}]
    accounts = LocalAuthenticator({EMAIL_A: {"user_id": USER_A, "password_hash": hash_password(PASSWORD)},
                                   EMAIL_B: {"user_id": USER_B, "password_hash": hash_password(PASSWORD)}})
    monkeypatch.setattr(auth_module, "_limiter", LoginRateLimiter())
    api.dependency_overrides[get_repository] = lambda: SupabaseRepository(fake.base_url, fake.key,
                                                                          "labextract-originals", http=fake)
    api.dependency_overrides[get_authenticator] = lambda: accounts
    yield fake
    api.dependency_overrides.clear()


def _signed_in(email):
    client = TestClient(app, base_url=BASE)
    assert client.post("/api/auth/login", json={"email": email, "password": PASSWORD},
                       headers={"Origin": BASE}).status_code == 200
    return client


def _pdf() -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    for n, text in enumerate(["Synthetic laboratory form", "First name: Dana", "Last name: Levi",
                              "Israeli ID: 000000018", "Sample received date: 23/08/2026",
                              "Sender: Synthetic Clinic", "Central number: 000987", "Commitment number: 000730001"]):
        page.insert_text((72, 80 + 24 * n), text, fontsize=11)
    data = doc.tobytes()
    doc.close()
    return data


def test_full_path_on_the_supabase_code_path(fake):
    client = _signed_in(EMAIL_A)
    token = client.cookies.get("lx_session")
    assert fake.tables["lx_sessions"][0]["token_hash"] == token_hash(token)       # only the hash is stored
    assert token not in str(fake.tables)

    origin = {"Origin": BASE}
    meta = client.post("/api/documents/upload", headers=origin,
                       files={"file": ("synthetic.pdf", _pdf(), "application/pdf")}).json()
    file_id = meta["file_id"]
    assert fake.object_types == {f"labextract-originals/demo/{file_id}.pdf": "application/pdf"}
    assert [row["workspace_id"] for row in fake.tables["lx_documents"]] == ["demo"]

    forms = client.post(f"/api/documents/{file_id}/drafts/init", headers=origin).json()
    form_id = forms[0]["form_id"]
    assert forms[0]["examinations"][0]["first_name"]["normalized"] == "Dana"
    assert client.get("/api/documents").json()[0]["processing_status"] == "DONE"

    sender = forms[0]["examinations"][0]["sender"]
    saved = client.patch(f"/api/review/{file_id}/{form_id}/field", headers=origin, json={
        "target": {"field": "sender"}, "value": "Synthetic Clinic North", "expected_version": sender["version"]})
    assert saved.status_code == 200
    assert fake.tables["lx_drafts"][0]["version"] == 2

    values = {"first_name": "Dana", "last_name": "Levi", "israeli_id": "000000018",
              "sample_received_date": "2026-08-23", "sender": "Synthetic Clinic North",
              "central_numbers": ["000987"], "commitment_numbers": ["000730001"]}
    assert client.put(f"/api/review/{file_id}/{form_id}", headers=origin,
                      json={**values, "approve": True}).status_code == 200
    excel = client.get(f"/api/documents/{file_id}/export.xlsx?confirm_partial=true")
    assert excel.status_code == 200
    row = [c.value for c in load_workbook(io.BytesIO(excel.content)).active[2]]
    assert row == ["Dana", "Levi", "000000018", "000987", "23/08/2026", "Synthetic Clinic North", "000730001"]
    assert client.get(f"/api/documents/{file_id}/file").content.startswith(b"%PDF")

    other = _signed_in(EMAIL_B)
    assert other.get("/api/documents").json() == []
    assert other.get(f"/api/documents/{file_id}/file").status_code == 404

    assert client.post("/api/auth/logout", headers=origin).status_code == 204
    assert fake.tables["lx_sessions"][0]["revoked_at"] is not None
    assert client.get(f"/api/documents/{file_id}/file").status_code == 401


def test_storage_failure_is_reported_not_hidden(fake):
    client = _signed_in(EMAIL_A)
    fake.fail_next["POST lx_documents"] = 503
    with pytest.raises(Exception):
        client.post("/api/documents/upload", headers={"Origin": BASE},
                    files={"file": ("synthetic.pdf", _pdf(), "application/pdf")})
    assert fake.objects == {} and fake.tables["lx_documents"] == []         # no orphan original
