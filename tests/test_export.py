import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app import app
from backend.excel_export import HEADERS, build_export
from extraction.mock_engine import make_mock_result
from shared.contracts import ApprovalStatus, ExtractionStatus, FieldName, FieldValue

client = TestClient(app)


def _sheet(forms):
    data, summary = build_export(forms)
    return load_workbook(io.BytesIO(data)).active, summary


def _is_orange(cell) -> bool:
    return cell.fill.fill_type == "solid" and cell.fill.start_color.rgb.endswith("FFC000")


def test_headers_and_one_row_per_commitment():
    ws, s = _sheet([make_mock_result().form])
    assert [c.value for c in ws[1]] == HEADERS
    assert ws.max_row == 1 + 3 and s.rows == 3          # three commitments, incl. the unreadable one
    assert ws.sheet_view.rightToLeft is True


def test_nothing_is_approved_so_everything_is_empty_and_orange():
    ws, s = _sheet([make_mock_result().form])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            assert c.value is None and _is_orange(c)
    assert s.is_partial and s.filename == "labextract_חלקי.xlsx" and s.orange_cells == 21


def test_approved_values_keep_leading_zeros_as_text():
    form = make_mock_result().form
    exam = form.examinations[0]
    ok = lambda fv: fv.model_copy(update={"approval_status": ApprovalStatus.HUMAN_APPROVED})
    exam.israeli_id = ok(exam.israeli_id)
    exam.central_numbers = [ok(x) for x in exam.central_numbers]
    exam.commitments[0].number = ok(exam.commitments[0].number)
    ws, _ = _sheet([form])
    assert ws["C2"].value == "000000018" and ws["C2"].data_type == "s" and not _is_orange(ws["C2"])
    assert ws["D2"].value == "000987;000988"
    assert ws["G2"].value == "000730001"
    assert ws["G3"].value is None and _is_orange(ws["G3"])      # partial 872?41596 is never exported


def test_formula_injection_is_neutralised():
    form = make_mock_result().form
    exam = form.examinations[0]
    evil = FieldValue(field_name=FieldName.FIRST_NAME, raw="=1+1", normalized="=1+1",
                      extraction_status=ExtractionStatus.COMPLETE, approval_status=ApprovalStatus.HUMAN_APPROVED)
    exam.first_name = evil
    ws, _ = _sheet([form])
    assert ws["A2"].value == "=1+1" and ws["A2"].data_type == "s"       # text, not 'f' (formula)


def test_confirmed_absent_is_empty_without_colour():
    form = make_mock_result().form
    exam = form.examinations[0]
    exam.sender = exam.sender.model_copy(update={"approval_status": ApprovalStatus.CONFIRMED_ABSENT})
    ws, _ = _sheet([form])
    assert ws["F2"].value is None and not _is_orange(ws["F2"])


def test_count_uncertain_creates_no_rows_but_warns():
    ws, s = _sheet([make_mock_result(scenario="count_unsure").form])
    assert ws.max_row == 1 and s.rows == 0 and s.uncounted_forms == 1 and s.forms_without_commitments == 1
    assert any("חסרות שורות" in w for w in s.warnings)


def test_failed_form_is_reported_not_exported():
    _, s = _sheet([make_mock_result(scenario="failed").form, make_mock_result(form_id="f2").form])
    assert s.failed_forms == 1 and s.rows == 3


# ---- API --------------------------------------------------------------------------------
def test_api_health_and_mock_result():
    assert client.get("/health").json() == {"status": "ok"}
    r = client.get("/mock/result")
    assert r.status_code == 200 and r.json()["form"]["examinations"][0]["commitments"][0]["number"]["normalized"] == "000730001"
    assert client.get("/mock/result?scenario=nope").status_code == 422


def test_partial_download_needs_explicit_confirmation():
    r = client.get("/mock/export.xlsx")
    assert r.status_code == 409 and r.json()["detail"]["needs_confirmation"] is True
    ok = client.get("/mock/export.xlsx?confirm_partial=true")
    assert ok.status_code == 200 and ok.content[:2] == b"PK"            # a real xlsx (zip) file
    assert "UTF-8''" in ok.headers["content-disposition"]


def test_summary_endpoint_and_multiple_forms():
    s = client.get("/mock/export/summary?scenarios=default,count_unsure,failed").json()
    assert s["rows"] == 3 and s["uncounted_forms"] == 1 and s["failed_forms"] == 1 and s["is_partial"] is True
