"""End to end through the REAL pipeline on this machine, A to Z, the way the site runs a document:
upload -> page by page (drafts/pages/N: local OCR + labx crops-first + automatic approval) -> review (fill / approve
one form like a reviewer) -> READY -> export summary + Excel. Same switches (backend/*.txt) and the server's
time limits; only the storage is local and the network is this machine's (no Vercel / Supabase).

python tools/local_e2e.py [--part 2] [--pages 1,2,3]
Calls Gemini (free keys only; answers are cached). Prints checks and counts, never patient values.
"""
import argparse
import io
import os
import sys
import tempfile
import time
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(1, str(ROOT / "tools"))
sys.stdout.reconfigure(encoding="utf-8")
# as the Dockerfile sets them on the server
os.environ.setdefault("LABEXTRACT_ENGINE", "hybrid")
os.environ.setdefault("LABEXTRACT_LABX", "on")
os.environ.setdefault("LABX_GEMINI_TIMEOUT_MS", "30000")
os.environ.setdefault("LABX_SLOT_BUDGET_S", "45")
os.environ.setdefault("LABX_PAGE_BUDGET_S", "150")
VERCEL_LIMIT_S = 300

import openpyxl  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

import measure_privacy as M  # noqa: E402
from backend.app import app  # noqa: E402  (loads the switches)
from backend.auth import Principal, require_principal  # noqa: E402
from backend.storage import DEFAULT_WORKSPACE, LocalRepository, get_repository  # noqa: E402
from shared.contracts import Form  # noqa: E402

failed = []


def check(name, ok, info=""):
    print(f"[{'PASS' if ok else 'FAIL'}] {name} {info}", flush=True)
    if not ok:
        failed.append(name)
    return ok


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--part", type=int, default=2, choices=sorted(M.PARTS))
    ap.add_argument("--pages", default="")
    a = ap.parse_args()
    pdf, T = M.PARTS[a.part][0], M.truth(a.part)
    folder = Path(tempfile.mkdtemp(prefix="labextract_e2e_"))
    app.dependency_overrides[get_repository] = lambda: LocalRepository(folder)
    app.dependency_overrides[require_principal] = lambda: Principal("e2e-reviewer", DEFAULT_WORKSPACE)  # signed in
    client = TestClient(app)
    print(f"switches: auto-approve={os.environ.get('LABX_AUTO_APPROVE')} mode={os.environ.get('LABX_MODE')}")

    # 1. upload: direct upload, as the site does for a file over 4 MB (start -> put the bytes -> complete)
    content = pdf.read_bytes()
    r = client.post("/documents/uploads", json={"filename": pdf.name, "size": len(content)})
    if not check("upload start", r.status_code == 200, f"{r.status_code} {len(content) // 1024} KB"):
        return
    started = r.json()
    r = client.put(f"/documents/uploads/{started['file_id']}/content", content=content)
    check("upload content", r.status_code == 204, str(r.status_code))
    r = client.post(f"/documents/uploads/{started['file_id']}/complete", json={"filename": pdf.name})
    if not check("upload complete", r.status_code == 200, str(r.status_code)):
        return
    meta = r.json()
    fid, count = meta["file_id"], meta["page_count"]
    check("page count", count == len(T), f"{count} pages, truth has {len(T)}")
    pages = [int(x) for x in a.pages.split(",")] if a.pages else list(range(1, count + 1))

    # 2. page by page, as the review screen drives it
    times = {}
    for p in pages:
        t0 = time.monotonic()
        r = client.post(f"/documents/{fid}/drafts/pages/{p}")
        times[p] = time.monotonic() - t0
        status = r.json()["form"]["processing_status"] if r.status_code == 200 else None
        check(f"page {p} processed", r.status_code == 200 and status == "DONE",
              f"{r.status_code} {status} {times[p]:.0f}s")
        check(f"page {p} inside the request limit", times[p] < VERCEL_LIMIT_S, f"{times[p]:.0f}s")

    # 3. what the reviewer gets: score every field against the truth, approvals included
    forms = [Form.model_validate(f) for f in client.get(f"/documents/{fid}/drafts").json()]
    tot = Counter()
    for f in forms:
        p = f.page_regions[0]["page"] + 1
        if p in pages:
            for field, marks in M.score(f, T[p]).items():
                if field not in ("first_name", "last_name"):
                    tot.update(marks)
    print(f"   fields: approved correct {tot['A']}, approved WRONG {tot['X']}, filled correct {tot['✓']}, "
          f"filled WRONG {tot['✗']}, review {tot['?']}, nothing {tot['·']}")
    check("no automatically approved value is wrong", tot["X"] == 0, str(tot["X"]))

    # 4. a reviewer closes one form: fill what is empty with the truth, approve every field -> READY
    f = forms[pages[0] - 1]
    t = T[pages[0]]
    exam = f.examinations[0]
    fills = {"first_name": "ישראל", "last_name": "ישראלי", "israeli_id": t["id"].zfill(9),
             "sample_received_date": t["date"], "sender": t["sender"]}
    for name, value in fills.items():
        fv = getattr(exam, name)
        if fv.approval_status.value not in ("AUTO_APPROVED", "HUMAN_APPROVED") and fv.normalized != value:
            r = client.patch(f"/review/{fid}/{f.form_id}/field",
                             json={"target": {"field": name}, "value": value, "expected_version": fv.version})
            check(f"reviewer fills {name}", r.status_code == 200, str(r.status_code))
    exam = Form.model_validate(client.get(f"/documents/{fid}/drafts").json()[pages[0] - 1]).examinations[0]
    targets = [{"field": k} for k in fills]
    targets += [{"field": "central_numbers", "index": i, "expected_value": c.normalized}
                for i, c in enumerate(exam.central_numbers) if c.normalized]
    targets += [{"field": "commitment_number", "commitment_id": c.commitment_id} for c in exam.commitments
                if c.number.normalized]
    for target in targets:
        cur = Form.model_validate(client.get(f"/documents/{fid}/drafts").json()[pages[0] - 1]).examinations[0]
        if target["field"] == "central_numbers":
            version = cur.central_numbers[target["index"]].version
        elif target["field"] == "commitment_number":
            version = next(c.number.version for c in cur.commitments if c.commitment_id == target["commitment_id"])
        else:
            version = getattr(cur, target["field"]).version
        r = client.post(f"/review/{fid}/{f.form_id}/field/approval",
                        json={"target": target, "action": "APPROVE", "expected_version": version})
        check(f"approve {target['field']}", r.status_code == 200, "" if r.status_code == 200 else r.text[:160])

    if any(i.reason.value == "COMMITMENT_COUNT_UNCERTAIN" for i in f.issues) or             any(c.existence_status.value == "UNCERTAIN" for c in exam.commitments):
        r = client.post(f"/review/{fid}/{f.form_id}/commitments/confirm",
                        json={"commitment_count": len(exam.commitments)})
        check("reviewer confirms no further commitments", r.status_code == 200, str(r.status_code))

    # 5. export
    r = client.get(f"/documents/{fid}/export/summary")
    check("export summary", r.status_code == 200, str(r.status_code))
    summary = r.json() if r.status_code == 200 else {}
    print("   summary:", {k: summary.get(k) for k in ("rows", "ready", "pending_rows", "orange_cells", "failed_forms")})
    check("the reviewed form is READY", (summary.get("ready") or 0) >= 1, str(summary.get("ready")))
    r = client.get(f"/documents/{fid}/export.xlsx?confirm_partial=true")
    if check("Excel download", r.status_code == 200 and r.content[:2] == b"PK", str(r.status_code)):
        ws = openpyxl.load_workbook(io.BytesIO(r.content)).worksheets[0]
        check("Excel has a row per commitment", ws.max_row - 1 == summary.get("rows"), f"{ws.max_row - 1} rows")
    print(f"page times: max {max(times.values()):.0f}s, mean {sum(times.values()) / len(times):.0f}s")
    print("RESULT:", "ALL PASSED" if not failed else f"{len(failed)} FAILED: {failed}")
    app.dependency_overrides.pop(get_repository, None)


if __name__ == "__main__":
    main()
